#!/usr/bin/env python3
import argparse
import json
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches


SLIDE_SIZES = {
    "16:9": (13.333, 7.5),
    "4:3": (10.0, 7.5),
}


def compose(plan_path: Path, image_paths: list[Path], output_path: Path) -> None:
    plan = json.loads(plan_path.read_text(encoding="utf-8"))
    aspect_ratio = plan.get("aspect_ratio", "16:9")
    if aspect_ratio not in SLIDE_SIZES:
        raise ValueError(f"Unsupported aspect ratio: {aspect_ratio}")
    if not image_paths:
        raise ValueError("At least one slide image is required")

    missing = [str(path) for path in image_paths if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing slide images: " + ", ".join(missing))

    width_in, height_in = SLIDE_SIZES[aspect_ratio]
    target_ratio = width_in / height_in
    presentation = Presentation()
    presentation.slide_width = Inches(width_in)
    presentation.slide_height = Inches(height_in)
    blank = presentation.slide_layouts[6]
    slide_specs = plan.get("slides", [])

    for index, image_path in enumerate(image_paths):
        slide = presentation.slides.add_slide(blank)
        with Image.open(image_path) as source:
            rgb = source.convert("RGB")
            crop_width = rgb.height * target_ratio
            crop_height = rgb.width / target_ratio
            if crop_width <= rgb.width:
                size = (round(crop_width), rgb.height)
            else:
                size = (rgb.width, round(crop_height))
            fitted = ImageOps.fit(rgb, size, method=Image.Resampling.LANCZOS)
            stream = BytesIO()
            fitted.save(stream, format="JPEG", quality=96, optimize=True)
            stream.seek(0)
            slide.shapes.add_picture(
                stream,
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        if index < len(slide_specs):
            spec = slide_specs[index]
            notes = []
            if spec.get("takeaway"):
                notes.append(f"Takeaway: {spec['takeaway']}")
            if spec.get("sources"):
                notes.append("Sources:")
                notes.extend(f"- {source}" for source in spec["sources"])
            if notes:
                slide.notes_slide.notes_text_frame.text = "\n".join(notes)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Compose slide images into a PPTX")
    parser.add_argument("--plan-file", type=Path, required=True)
    parser.add_argument("--slide-images", type=Path, nargs="+", required=True)
    parser.add_argument("--output-file", type=Path, required=True)
    args = parser.parse_args()
    compose(args.plan_file, args.slide_images, args.output_file)
    print(f"Created {args.output_file} with {len(args.slide_images)} slides")


if __name__ == "__main__":
    main()
